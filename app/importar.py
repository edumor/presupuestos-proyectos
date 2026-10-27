# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
"""Importación desde Excel: equipo, gastos fijos, servicios, clientes, horas, facturas y cobros.

Se usa la planilla modelo (plantilla()). La vista previa y la importación ejecutan exactamente el
mismo proceso; la vista previa deshace todo al final. Si hay algún error no se importa nada.
Los registros se buscan por nombre: si ya existen se actualizan (equipo, gastos, servicios,
clientes, facturas); las horas y los cobros que ya estén cargados iguales se omiten.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from . import CREDITO, costeo, facturacion
from .models import (Cliente, Cobro, CostoFijo, Etapa, Factura, HoraReal, Periodicidad, Persona, Proyecto, Rol,
                     Servicio, Tarea, TipoPersona)
from .schemas import (MONEDAS, ClienteIn, CobroIn, CostoFijoIn, FacturaIn, HoraIn, PersonaIn, ServicioIn)


class ErrorFila(Exception):
    pass


# hoja: [(columna, obligatoria, ayuda)]
HOJAS: dict[str, list[tuple[str, bool, str]]] = {
    "Equipo": [("Nombre", True, "Nombre de la persona (si ya existe, se actualiza)"),
               ("Rol", True, "Si el rol no existe, se crea"),
               ("Tipo", True, "Empleado o Autónomo"),
               ("Moneda", False, "Moneda del sueldo o tarifa (vacío = moneda base)"),
               ("Sueldo bruto anual", False, "Empleados: 13 sueldos (incluye aguinaldo)"),
               ("Tarifa por hora", False, "Autónomos"),
               ("Contribuciones %", False, "Empleados: por ejemplo 28,4 (vacío = el de Configuración)"),
               ("Activo", False, "Sí o No (vacío = Sí)")],
    "Gastos fijos": [("Concepto", True, "Si ya existe, se actualiza"),
                     ("Periodicidad", True, "Mensual o Anual"),
                     ("Moneda", False, "Vacío = moneda base"),
                     ("Importe", True, "")],
    "Servicios": [("Nombre", True, "Si ya existe, se actualiza"), ("Descripción", False, ""),
                  ("Rol sugerido", False, "Debe existir (o venir en la hoja Equipo)"),
                  ("Horas sugeridas", False, ""), ("% sobre el total", False, "Por ejemplo 10 para gestión"),
                  ("Etapa sugerida", False, "Inicio, Planificación, Ejecución o Cierre")],
    "Clientes": [("Nombre", True, "Si ya existe, se actualiza"), ("País", False, ""), ("Contacto", False, ""),
                 ("Moneda a mostrar", False, "Moneda del equivalente en las propuestas (por ejemplo USD)")],
    "Horas": [("Fecha", True, ""), ("Persona", True, "Debe existir"), ("Proyecto", True, "Debe existir"),
              ("Tarea", True, "Tarea del proyecto (debe existir)"), ("Horas", True, ""), ("Nota", False, "")],
    "Facturas": [("Proyecto", True, "Debe existir"), ("Número", True, "Si ya existe en el proyecto, se actualiza"),
                 ("Fecha", True, ""), ("Vencimiento", False, ""),
                 ("Moneda", False, "Vacío = moneda del contrato"), ("Importe neto", True, "Sin impuestos"),
                 ("Impuestos", False, "IVA u otros"), ("Cotización", False, "Vacío = la del día"),
                 ("Concepto", False, ""), ("Etapa", False, "Nombre de la etapa del proyecto"),
                 ("Anulada", False, "Sí o No")],
    "Cobros": [("Proyecto", True, ""), ("Número de factura", True, "La factura debe existir o venir en la hoja Facturas"),
               ("Fecha", True, ""), ("Importe", True, "En la moneda de la factura"),
               ("Cotización", False, "Vacío = la del día"), ("Notas", False, "")],
}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("%", " pct ").split())


def _texto(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _numero(v, campo: str) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    for simbolo in ("$", *MONEDAS):
        s = s.replace(simbolo, "")
    if "," in s and "." in s:  # el último separador es el decimal: 1.234,56 o 1,234.56
        miles = "." if s.rfind(",") > s.rfind(".") else ","
        s = s.replace(miles, "").replace(",", ".")
    elif s.count(",") == 1:  # 4,5
        s = s.replace(",", ".")
    elif s.count(",") > 1 or s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", s):  # 26.000.000 o 1.500
        s = s.replace(",", "").replace(".", "")
    try:
        return float(s)
    except ValueError:
        raise ErrorFila(f"{campo}: «{v}» no es un número") from None


def _fecha(v, campo: str) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ErrorFila(f"{campo}: «{v}» no es una fecha (use dd/mm/aaaa)")


def _si_no(v, defecto: bool) -> bool:
    t = _norm(v)
    if not t:
        return defecto
    if t in ("si", "s", "true", "1", "x", "yes"):
        return True
    if t in ("no", "n", "false", "0"):
        return False
    raise ErrorFila(f"«{v}» debe ser Sí o No")


def _moneda(db: Session, v) -> str:
    m = (_texto(v) or costeo.moneda_base(db)).upper()
    if m not in MONEDAS:
        raise ErrorFila(f"Moneda «{m}» no admitida (use {', '.join(MONEDAS)})")
    return m


def _validacion(e: ValidationError) -> str:
    return "; ".join(str(x["msg"]).removeprefix("Value error, ") for x in e.errors())


def _por_nombre(db: Session, modelo, nombre: str, **filtro):
    q = select(modelo).where(func.lower(modelo.nombre) == nombre.lower())
    for k, v in filtro.items():
        q = q.where(getattr(modelo, k) == v)
    return db.scalars(q).first()


def _proyecto(db: Session, nombre: str) -> Proyecto:
    p = _por_nombre(db, Proyecto, nombre)
    if p is None:
        raise ErrorFila(f"No existe el proyecto «{nombre}»")
    return p


# --- procesamiento de cada hoja: devuelve ("nuevo" | "actualizado" | "omitido", descripción) ---
def _equipo(db: Session, f: dict):
    nombre, rol_nom = _texto(f["nombre"]), _texto(f["rol"])
    tipo_txt = _norm(f["tipo"])
    tipo = TipoPersona.empleado if tipo_txt.startswith("emp") else TipoPersona.autonomo if tipo_txt.startswith("aut") else None
    if tipo is None:
        raise ErrorFila("Tipo debe ser Empleado o Autónomo")
    rol = _por_nombre(db, Rol, rol_nom)
    if rol is None:
        rol = Rol(nombre=rol_nom)
        db.add(rol); db.flush()
    ss = _numero(f.get("contribuciones pct"), "Contribuciones %")
    try:
        d = PersonaIn(nombre=nombre, rol_id=rol.id, tipo=tipo, moneda=_moneda(db, f.get("moneda")),
                      sueldo_bruto_anual=_numero(f.get("sueldo bruto anual"), "Sueldo") if tipo == TipoPersona.empleado else None,
                      tarifa_hora=_numero(f.get("tarifa por hora"), "Tarifa") if tipo == TipoPersona.autonomo else None,
                      ss_empresa=(ss / 100 if ss and ss > 1 else ss) if tipo == TipoPersona.empleado else None,
                      activo=_si_no(f.get("activo"), True))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    p = _por_nombre(db, Persona, nombre)
    accion = "actualizado" if p else "nuevo"
    if p is None:
        p = Persona()
    for k, v in d.model_dump().items():
        setattr(p, k, v)
    db.add(p); db.flush()
    importe = d.sueldo_bruto_anual or d.tarifa_hora
    return accion, f"{nombre} · {rol_nom} · {tipo.value} · {d.moneda} {importe:,.2f}"


def _gasto(db: Session, f: dict):
    concepto = _texto(f["concepto"])
    per = _norm(f["periodicidad"])
    periodicidad = Periodicidad.mensual if per.startswith("men") else Periodicidad.anual if per.startswith("anu") else None
    if periodicidad is None:
        raise ErrorFila("Periodicidad debe ser Mensual o Anual")
    try:
        d = CostoFijoIn(concepto=concepto, moneda=_moneda(db, f.get("moneda")),
                        importe=_numero(f["importe"], "Importe"), periodicidad=periodicidad)
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    c = db.scalars(select(CostoFijo).where(func.lower(CostoFijo.concepto) == concepto.lower())).first()
    accion = "actualizado" if c else "nuevo"
    if c is None:
        c = CostoFijo()
    for k, v in d.model_dump(exclude={"vigente_desde", "vigente_hasta"}).items():
        setattr(c, k, v)
    db.add(c); db.flush()
    return accion, f"{concepto} · {d.moneda} {d.importe:,.2f} {periodicidad.value}"


def _servicio(db: Session, f: dict):
    nombre = _texto(f["nombre"])
    rol_id = None
    if _texto(f.get("rol sugerido")):
        rol = _por_nombre(db, Rol, _texto(f["rol sugerido"]))
        if rol is None:
            raise ErrorFila(f"No existe el rol «{_texto(f['rol sugerido'])}»")
        rol_id = rol.id
    pct = _numero(f.get("pct sobre el total"), "% sobre el total")
    try:
        d = ServicioIn(nombre=nombre, descripcion=_texto(f.get("descripcion")), rol_sugerido_id=rol_id,
                       horas_sugeridas=_numero(f.get("horas sugeridas"), "Horas sugeridas"),
                       pct_sobre_total=(pct / 100 if pct and pct > 1 else pct),
                       etapa_sugerida=_texto(f.get("etapa sugerida")))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    s = _por_nombre(db, Servicio, nombre)
    accion = "actualizado" if s else "nuevo"
    if s is None:
        s = Servicio()
    for k, v in d.model_dump().items():
        setattr(s, k, v)
    db.add(s); db.flush()
    return accion, nombre


def _cliente(db: Session, f: dict):
    nombre = _texto(f["nombre"])
    try:
        d = ClienteIn(nombre=nombre, pais=_texto(f.get("pais")), contacto=_texto(f.get("contacto")),
                      moneda_mostrar=_moneda(db, f.get("moneda a mostrar")))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    c = _por_nombre(db, Cliente, nombre)
    accion = "actualizado" if c else "nuevo"
    if c is None:
        c = Cliente()
    for k, v in d.model_dump().items():
        setattr(c, k, v)
    db.add(c); db.flush()
    return accion, nombre


def _hora(db: Session, f: dict):
    fecha = _fecha(f["fecha"], "Fecha")
    per = _por_nombre(db, Persona, _texto(f["persona"]))
    if per is None:
        raise ErrorFila(f"No existe la persona «{_texto(f['persona'])}»")
    p = _proyecto(db, _texto(f["proyecto"]))
    t = _por_nombre(db, Tarea, _texto(f["tarea"]), proyecto_id=p.id)
    if t is None:
        raise ErrorFila(f"El proyecto «{p.nombre}» no tiene la tarea «{_texto(f['tarea'])}»")
    try:
        d = HoraIn(tarea_id=t.id, persona_id=per.id, fecha=fecha, horas=_numero(f["horas"], "Horas"),
                   nota=_texto(f.get("nota")))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    desc = f"{fecha:%d/%m/%Y} · {per.nombre} · {p.nombre} / {t.nombre} · {d.horas:g} h"
    igual = db.scalars(select(HoraReal).where(HoraReal.tarea_id == t.id, HoraReal.persona_id == per.id,
                                              HoraReal.fecha == fecha, HoraReal.horas == d.horas)).first()
    if igual:
        return "omitido", desc + " (ya estaba cargada)"
    h = HoraReal(**d.model_dump())
    h.costo_hora_base = round(costeo.costo_hora(db, per, fecha)["costo_total_base"], 2)
    db.add(h); db.flush()
    return "nuevo", desc


def _factura(db: Session, f: dict):
    from .routers.facturacion import _guardar
    p = _proyecto(db, _texto(f["proyecto"]))
    numero = _texto(f["numero"])
    etapa_id = None
    if _texto(f.get("etapa")):
        e = _por_nombre(db, Etapa, _texto(f["etapa"]), proyecto_id=p.id)
        if e is None:
            raise ErrorFila(f"El proyecto no tiene la etapa «{_texto(f['etapa'])}»")
        etapa_id = e.id
    try:
        d = FacturaIn(numero=numero, concepto=_texto(f.get("concepto")), etapa_id=etapa_id,
                      fecha=_fecha(f["fecha"], "Fecha"), vencimiento=_fecha(f.get("vencimiento"), "Vencimiento"),
                      moneda=_texto(f.get("moneda")), importe_neto=_numero(f["importe neto"], "Importe neto"),
                      impuesto=_numero(f.get("impuestos"), "Impuestos") or 0,
                      cotizacion=_numero(f.get("cotizacion"), "Cotización"), anulada=_si_no(f.get("anulada"), False))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    fa = db.scalars(select(Factura).where(Factura.proyecto_id == p.id, Factura.numero == numero)).first()
    accion = "actualizado" if fa else "nuevo"
    if fa is None:
        fa = Factura(proyecto_id=p.id)
        fa.cobros = []
    try:
        _guardar(db, fa, p, d)
    except HTTPException as e:
        raise ErrorFila(e.detail) from None
    db.add(fa); db.flush()
    return accion, f"{p.nombre} · {numero} · {fa.moneda} {d.importe_neto:,.2f}"


def _cobro(db: Session, f: dict):
    from .routers.facturacion import _cotizacion
    p = _proyecto(db, _texto(f["proyecto"]))
    numero = _texto(f["numero de factura"])
    fa = db.scalars(select(Factura).where(Factura.proyecto_id == p.id, Factura.numero == numero)).first()
    if fa is None:
        raise ErrorFila(f"No existe la factura «{numero}» en el proyecto «{p.nombre}»")
    try:
        d = CobroIn(fecha=_fecha(f["fecha"], "Fecha"), importe=_numero(f["importe"], "Importe"),
                    cotizacion=_numero(f.get("cotizacion"), "Cotización"), notas=_texto(f.get("notas")))
    except ValidationError as e:
        raise ErrorFila(_validacion(e)) from None
    desc = f"{p.nombre} · {numero} · {d.fecha:%d/%m/%Y} · {fa.moneda} {d.importe:,.2f}"
    if any(c.fecha == d.fecha and abs(float(c.importe) - d.importe) < 0.005 for c in fa.cobros):
        return "omitido", desc + " (ya estaba cargado)"
    if fa.anulada:
        raise ErrorFila("La factura está anulada")
    if d.fecha < fa.fecha:
        raise ErrorFila("El cobro es anterior a la factura")
    saldo = facturacion.total(fa) - facturacion.cobrado(fa)
    if d.importe > saldo + facturacion.TOLERANCIA:
        raise ErrorFila(f"El cobro supera el saldo de la factura ({fa.moneda} {saldo:,.2f})")
    cot = _cotizacion(db, fa.moneda, d.fecha, d.cotizacion)
    fa.cobros.append(Cobro(fecha=d.fecha, importe=d.importe, cotizacion=cot, importe_base=round(d.importe * cot, 2),
                           notas=d.notas))
    db.flush()
    return "nuevo", desc


PROCESOS = {"Equipo": _equipo, "Gastos fijos": _gasto, "Servicios": _servicio, "Clientes": _cliente,
            "Horas": _hora, "Facturas": _factura, "Cobros": _cobro}


def _filas(ws, columnas: list[tuple[str, bool, str]]):
    """Encabezado en la primera fila no vacía; devuelve (número de fila, {columna normalizada: valor})."""
    filas = ws.iter_rows(values_only=True)
    enc, n = None, 0
    for fila in filas:
        n += 1
        if any(v is not None and str(v).strip() for v in fila):
            enc = [_norm(v).rstrip(" *").strip() for v in fila]
            break
    if enc is None:
        return
    faltan = [c for c, oblig, _ in columnas if oblig and _norm(c) not in enc]
    if faltan:
        raise ErrorFila("Faltan columnas: " + ", ".join(faltan))
    for fila in filas:
        n += 1
        if not any(v is not None and str(v).strip() for v in fila):
            continue
        yield n, {k: v for k, v in zip(enc, fila) if k}


def procesar(db: Session, contenido: bytes, aplicar: bool) -> dict:
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(422, "El archivo no es un Excel (.xlsx) válido.") from None
    costeo.config(db)  # crea la configuración si falta (hace commit): antes de tocar nada
    nombres = {_norm(n): n for n in wb.sheetnames}
    hojas, errores = [], 0
    try:
        for hoja, columnas in HOJAS.items():
            real = nombres.get(_norm(hoja))
            if real is None:
                continue
            r = dict(hoja=hoja, nuevos=0, actualizados=0, omitidos=0, errores=[], filas=[])
            try:
                for n, f in _filas(wb[real], columnas):
                    try:
                        for c, oblig, _ in columnas:
                            if oblig and _texto(f.get(_norm(c))) is None:
                                raise ErrorFila(f"Falta «{c}»")
                        # sin SAVEPOINT (pysqlite los confirma solos): cada proceso valida todo antes de
                        # agregar a la sesión, y si hubo algún error al final se deshace todo
                        accion, desc = PROCESOS[hoja](db, f)
                        r[{"nuevo": "nuevos", "actualizado": "actualizados", "omitido": "omitidos"}[accion]] += 1
                        r["filas"].append(dict(fila=n, accion=accion, detalle=desc))
                    except ErrorFila as e:
                        r["errores"].append(dict(fila=n, mensaje=str(e)))
                    except costeo.FaltaCotizacion as e:
                        r["errores"].append(dict(fila=n, mensaje=f"Falta la cotización de {e.moneda} para esa fecha"))
            except ErrorFila as e:
                r["errores"].append(dict(fila=1, mensaje=str(e)))
            errores += len(r["errores"])
            if r["filas"] or r["errores"]:
                hojas.append(r)
        if not hojas:
            raise HTTPException(422, "No encontré datos: usá la planilla modelo (hojas Equipo, Gastos fijos, "
                                     "Servicios, Clientes, Horas, Facturas, Cobros).")
        importado = aplicar and errores == 0
        if importado:
            db.commit()
        else:
            db.rollback()
    except BaseException:
        db.rollback()
        raise
    return dict(importado=importado, errores=errores, hojas=hojas,
                total=sum(h["nuevos"] + h["actualizados"] for h in hojas))


def plantilla(db: Session) -> bytes:
    """Planilla modelo: una hoja por tipo de dato, con instrucciones y listas desplegables."""
    base = costeo.moneda_base(db)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    ins["A1"] = "Importar datos a Presupuestos"
    ins["A1"].font = Font(bold=True, size=14)
    lineas = [
        "Completá solo las hojas que necesites; las demás pueden quedar vacías o borrarse.",
        "No cambies los títulos de la primera fila. Las columnas con * son obligatorias.",
        f"Los importes van en la moneda indicada (vacío = {base}). Fechas en formato dd/mm/aaaa.",
        "Equipo, gastos fijos, servicios, clientes y facturas: si el nombre (o número) ya existe, se actualiza.",
        "Horas y cobros: si ya hay uno igual cargado, se omite. Los proyectos y tareas deben existir.",
        "Orden de importación: Equipo, Gastos fijos, Servicios, Clientes, Horas, Facturas, Cobros.",
        "En el programa: Datos → Importar desde Excel. Primero vas a ver una vista previa; si hay errores,"
        " no se importa nada hasta corregirlos.",
    ]
    for i, t in enumerate(lineas, 3):
        ins.cell(i, 1, "• " + t)
    fila = len(lineas) + 5
    for hoja, columnas in HOJAS.items():
        ins.cell(fila, 1, hoja).font = Font(bold=True)
        fila += 1
        for c, oblig, ayuda in columnas:
            ins.cell(fila, 1, ("  " + c + (" *" if oblig else "")))
            ins.cell(fila, 2, ayuda).font = Font(color="5D6677")
            fila += 1
        fila += 1
    ins.cell(fila + 1, 1, CREDITO).font = Font(italic=True, size=8, color="8A93A3")
    ins.column_dimensions["A"].width = 34
    ins.column_dimensions["B"].width = 80

    enc = PatternFill("solid", fgColor="2456C8")
    listas = {"Tipo": '"Empleado,Autónomo"', "Periodicidad": '"Mensual,Anual"', "Activo": '"Sí,No"',
              "Anulada": '"Sí,No"', "Moneda": '"' + ",".join(MONEDAS) + '"', "Moneda a mostrar": '"' + ",".join(MONEDAS) + '"',
              "Etapa sugerida": '"Inicio,Planificación,Ejecución,Cierre"'}
    for hoja, columnas in HOJAS.items():
        ws = wb.create_sheet(hoja)
        for j, (c, oblig, ayuda) in enumerate(columnas, 1):
            cel = ws.cell(1, j, c + (" *" if oblig else ""))
            cel.fill, cel.font = enc, Font(bold=True, color="FFFFFF")
            cel.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[cel.column_letter].width = max(14, len(c) + 6)
            if c in listas:
                dv = DataValidation(type="list", formula1=listas[c], allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{cel.column_letter}2:{cel.column_letter}1000")
            if c in ("Fecha", "Vencimiento"):
                for i in range(2, 1001):
                    ws.cell(i, j).number_format = "dd/mm/yyyy"
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
