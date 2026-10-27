# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
"""Excel para el cliente (español o inglés): propuesta o informe de avance, sin costos internos."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font

from . import excel, i18n
from .excel import FMT_NUM, FMT_PCT, fmt_moneda


def _cabecera(ws, d: dict, titulo: str) -> int:
    """Encabezado con los datos de la empresa (el logo va solo en la barra de los HTML)."""
    t, emp = d["t"], d["empresa"]
    fila = 1
    if emp["nombre"]:
        ws.cell(fila, 1, emp["nombre"]).font = Font(bold=True, color="1C2230")
        fila += 1
    for linea in emp["lineas"]:
        ws.cell(fila, 1, linea).font = Font(color="5D6677")
        fila += 1
    fila += 1
    ws.cell(fila, 1, titulo).font = Font(bold=True, size=16)
    fila += 2
    p = d["proyecto"]
    return excel.ficha(ws, fila, [(t["cliente"], p["cliente"], None), (t["proyecto"], p["nombre"], None),
                                  (t["fecha"], d["fecha"], None), (t["referencia"], d["referencia"], None),
                                  (t["inicio_proy"], p["inicio"], None), (t["fin_proy"], p["fin"], None)]) + 1


def _pie(ws, fila: int, d: dict) -> None:
    ws.cell(fila + 1, 1, d["credito"]).font = Font(italic=True, size=8, color="8A93A3")


def propuesta(d: dict) -> bytes:
    t, mb, fx = d["t"], d["moneda"], d["fx"]
    wb = excel.libro()
    ws = wb.create_sheet(t["hoja_propuesta"])
    fila = _cabecera(ws, d, t["t_propuesta"])
    fila = excel.ficha(ws, fila, [(t["esfuerzo"], d["horas_totales"], FMT_NUM), (t["valido"], d["valido_hasta"], None)])
    if d["proyecto"]["alcance"]:
        ws.cell(fila + 1, 1, t["alcance"]).font = Font(bold=True, size=12)
        c = ws.cell(fila + 2, 1, d["proyecto"]["alcance"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila + 2, start_column=1, end_row=fila + 2, end_column=6)
        ws.row_dimensions[fila + 2].height = max(30, 15 * (len(d["proyecto"]["alcance"]) // 90 + 1))
        fila += 4
    fb = fmt_moneda(mb)
    fo = fmt_moneda(fx["moneda"]) if fx else None
    cols = [(t["etapa"], None), (t["inicio"], None), (t["fin"], None), (t["hito"], None), (t["horas_est"], FMT_NUM),
            (f"{t['precio']} ({mb})", fb)]
    if fx:
        cols.append((f"{t['precio']} ({fx['moneda']})", fo))
    datos = [[e["nombre"], e["inicio"], e["fin"], e["hito"], e["horas"], e["precio"]]
             + ([round(e["precio"] * fx["factor"], 2)] if fx else []) for e in d["etapas"]]
    enc, fin = excel.tabla(ws, fila + 1, cols, datos, titulo=t["inversion_etapa"],
                           anchos={1: 34, 2: 13, 3: 13, 4: 34, 5: 12, 6: 18, 7: 18})
    filas_tot = []
    if d["impuesto_pct"]:
        filas_tot += [(t["subtotal"], d["precio"]),
                      (f"{t['impuesto']} ({i18n.numero(d['impuesto_pct'] * 100, 2, d['idioma'])}%)", d["impuesto"])]
    filas_tot.append((t["total"], d["total"]))
    for etiqueta, v in filas_tot:
        fin += 1
        ws.cell(fin, 1, etiqueta).font = Font(bold=True)
        c = ws.cell(fin, 6, v)
        c.number_format, c.font = fb, Font(bold=True)
        if fx:
            c2 = ws.cell(fin, 7, round(v * fx["factor"], 2))
            c2.number_format, c2.font = fo, Font(bold=True)
    fin += 2
    ws.cell(fin, 1, t["nota_reparto"]).font = Font(italic=True, color="5D6677", size=9)
    if fx:
        fin += 1
        ws.cell(fin, 1, i18n.nota_fx(fx, mb, t, d["idioma"])).font = Font(italic=True, color="5D6677", size=9)
    if d.get("nota_ajuste"):
        fin += 1
        ws.cell(fin, 1, d["nota_ajuste"]).font = Font(italic=True, color="5D6677", size=9)
    if d["condiciones"]:
        fin += 2
        ws.cell(fin, 1, t["condiciones"]).font = Font(bold=True, size=12)
        c = ws.cell(fin + 1, 1, d["condiciones"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fin + 1, start_column=1, end_row=fin + 1, end_column=7)
        ws.row_dimensions[fin + 1].height = max(30, 15 * (d["condiciones"].count("\n") + len(d["condiciones"]) // 100 + 1))
        fin += 1
    _pie(ws, fin + 1, d)

    wa = wb.create_sheet(t["hoja_actividades"])
    cols = [(t["etapa"], None), (t["actividad"], None), (t["inicio"], None), (t["fin"], None)]
    filas = [[e["nombre"], ta["nombre"], ta["inicio"], ta["fin"]] for e in d["etapas"] for ta in e["tareas"]]
    enc, fin = excel.tabla(wa, 1, cols, filas, anchos={1: 26, 2: 44, 3: 13, 4: 13})
    wa.freeze_panes = "A2"
    _pie(wa, fin + 1, d)
    return excel.a_bytes(wb)


def informe_avance(d: dict) -> bytes:
    t = d["t"]
    ne = len(d["etapas"])
    wb = excel.libro()
    ws = wb.create_sheet(t["hoja_avance"])
    fila = _cabecera(ws, d, t["t_avance"])
    fila = excel.ficha(ws, fila, [(t["avance_general"], d["avance"], FMT_PCT),
                                  (t["plan_fecha"], d["avance_planificado"], FMT_PCT),
                                  (t["etapas_term"], f"{d['etapas_terminadas']} {t['de']} {ne}", None),
                                  (t["etapas_aprob_cli"], f"{d['etapas_aprobadas']} {t['de']} {ne}", None)])
    if d["proximo_hito"]:
        ph = d["proximo_hito"]
        fila = excel.ficha(ws, fila, [(t["proximo_hito"], f"{ph['hito']} ({ph['etapa']})", None),
                                      (t["fecha_prevista"], ph["fecha"], None)])
    cols = [(t["etapa"], None), (t["inicio_previsto"], None), (t["fin_previsto"], None), (t["avance"], FMT_PCT),
            (t["planificado"], FMT_PCT), (t["estado"], None), (t["hito"], None), (t["hito_entregado"], None),
            (t["aprobada"], None), (t["fecha_aprob"], None)]
    datos = [[e["nombre"], e["inicio"], e["fin"], e["avance"], e["avance_planificado"], e["estado"], e["hito"],
              e["hito_cumplido"], t["si"] if e["aprobada"] else t["no"], e["fecha_aprobacion"]] for e in d["etapas"]]
    enc, fin = excel.tabla(ws, fila + 1, cols, datos, titulo=t["avance_etapa"],
                           anchos={1: 30, 2: 13, 3: 13, 4: 11, 5: 11, 6: 16, 7: 30, 8: 14, 9: 12, 10: 13})
    if datos:
        excel.grafico_barras(ws, t["graf_avance"], enc, fin, 1, [4, 5], f"A{fin + 3}", horizontal=True, eje_y="%")
    _pie(ws, fin + 20, d)

    wa = wb.create_sheet(t["hoja_actividades"])
    cols = [(t["etapa"], None), (t["actividad"], None), (t["inicio_previsto"], None), (t["fin_previsto"], None),
            (t["avance"], FMT_PCT)]
    filas = [[e["nombre"], ta["nombre"], ta["inicio"], ta["fin"], ta["avance"]] for e in d["etapas"] for ta in e["tareas"]]
    enc, fin = excel.tabla(wa, 1, cols, filas, anchos={1: 26, 2: 44, 3: 13, 4: 13, 5: 11})
    wa.freeze_panes = "A2"
    _pie(wa, fin + 1, d)
    return excel.a_bytes(wb)
