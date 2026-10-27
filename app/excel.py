# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
"""Exportación a Excel (.xlsx) con formato, filtros y gráficos nativos."""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

AZUL = "2456C8"
GRIS = "F2F4F7"
ENCABEZADO = PatternFill("solid", fgColor="1C2230")
FUENTE_ENC = Font(color="FFFFFF", bold=True)
LINEA = Border(bottom=Side(style="thin", color="E3E6EB"))

FMT_BASE = "#,##0.00"
FMT_PCT = "0.0%"
FMT_NUM = "#,##0.00"
FMT_FECHA = "dd/mm/yyyy"


def fmt_moneda(moneda: str) -> str:
    """Formato numérico de Excel con el código de moneda delante, p. ej. \"ARS\" 1.234,56."""
    return f'"{moneda}" #,##0.00'


def libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def a_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valor(v):
    if isinstance(v, Decimal):
        return float(v)
    if hasattr(v, "value") and not isinstance(v, (str, int, float)):  # Enum
        return v.value
    return v


def tabla(ws: Worksheet, fila: int, columnas: list[tuple[str, str | None]], datos: list[list],
          titulo: str | None = None, anchos: dict[int, int] | None = None) -> tuple[int, int]:
    """Escribe una tabla con encabezado. columnas = [(nombre, formato)]. Devuelve (fila_enc, fila_fin)."""
    if titulo:
        ws.cell(fila, 1, titulo).font = Font(bold=True, size=12)
        fila += 1
    for j, (nombre, _) in enumerate(columnas, 1):
        c = ws.cell(fila, j, nombre)
        c.fill, c.font = ENCABEZADO, FUENTE_ENC
        c.alignment = Alignment(vertical="center", wrap_text=True)
    enc = fila
    for reg in datos:
        fila += 1
        for j, ((_, fmt), v) in enumerate(zip(columnas, reg), 1):
            c = ws.cell(fila, j, _valor(v))
            c.border = LINEA
            if fmt:
                c.number_format = fmt
            elif isinstance(v, (date, datetime)):
                c.number_format = FMT_FECHA
    for j, (nombre, fmt) in enumerate(columnas, 1):
        ancho = (anchos or {}).get(j) or max(12, min(45, len(nombre) + 4,
                                                     max((len(str(_valor(r[j - 1]) or "")) + 2 for r in datos), default=12)))
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, ancho)
    return enc, fila


def hoja_datos(wb: Workbook, nombre: str, columnas, datos, filtro=True) -> Worksheet:
    ws = wb.create_sheet(nombre[:31])
    enc, fin = tabla(ws, 1, columnas, datos)
    ws.freeze_panes = "A2"
    if filtro and datos:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{fin}"
    return ws


def ficha(ws: Worksheet, fila: int, pares: list[tuple[str, object, str | None]]) -> int:
    for etiqueta, valor, fmt in pares:
        a = ws.cell(fila, 1, etiqueta)
        a.font = Font(color="5D6677")
        b = ws.cell(fila, 2, _valor(valor))
        b.font = Font(bold=True)
        if fmt:
            b.number_format = fmt
        elif isinstance(valor, (date, datetime)):
            b.number_format = FMT_FECHA
        fila += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 30)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 22)
    return fila


def grafico_lineas(ws: Worksheet, titulo: str, fila_enc: int, fila_fin: int, col_x: int, cols_y: list[int],
                   ancla: str, eje_y: str = "") -> None:
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = titulo, eje_y, 8, 18
    for c in cols_y:
        ch.add_data(Reference(ws, min_col=c, min_row=fila_enc, max_row=fila_fin), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=col_x, min_row=fila_enc + 1, max_row=fila_fin))
    ch.x_axis.number_format = "dd/mm"
    ws.add_chart(ch, ancla)


def grafico_barras(ws: Worksheet, titulo: str, fila_enc: int, fila_fin: int, col_x: int, cols_y: list[int],
                   ancla: str, horizontal: bool = False, eje_y: str = "") -> None:
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.y_axis.title, ch.height, ch.width = titulo, eje_y, 8, 18
    for c in cols_y:
        ch.add_data(Reference(ws, min_col=c, min_row=fila_enc, max_row=fila_fin), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=col_x, min_row=fila_enc + 1, max_row=fila_fin))
    ws.add_chart(ch, ancla)
