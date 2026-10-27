# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
"""Contratos en dólares, ajuste por IPC, aumento de costos y consultas al BCRA (simuladas)."""
from datetime import date

import pytest

from app import bcra, contratos


@pytest.fixture()
def proyecto(client):
    cli = client.post("/api/clientes", json={"nombre": "C"}).json()
    client.post("/api/cotizaciones", json={"fecha": "2026-01-01", "moneda": "USD", "valor": 1000})
    client.post("/api/cotizaciones", json={"fecha": "2026-06-01", "moneda": "USD", "valor": 1500})
    p = client.post("/api/proyectos", json={"cliente_id": cli["id"], "nombre": "P", "fecha_inicio": "2026-01-05",
                                            "fecha_fin": "2026-12-20", "precio_base": 1_000_000}).json()
    return client, p


def test_contrato_en_dolares_se_reexpresa(proyecto):
    client, p = proyecto
    body = {**p, "moneda_contrato": "USD", "precio_contrato": 1000}
    out = client.put(f"/api/proyectos/{p['id']}", json=body).json()
    assert out["moneda_contrato"] == "USD" and out["precio_base"] == pytest.approx(1_500_000)
    s = client.get(f"/api/proyectos/{p['id']}/seguimiento", params={"fecha": "2026-03-01"}).json()
    assert s["precio_base"] == pytest.approx(1_000_000)  # a la cotización de enero
    s = client.get(f"/api/proyectos/{p['id']}/seguimiento", params={"fecha": "2026-07-01"}).json()
    assert s["precio_base"] == pytest.approx(1_500_000) and s["contrato"]["tipo"] == "moneda"
    # el ajuste por IPC no se combina con un contrato en otra moneda
    assert client.put(f"/api/proyectos/{p['id']}", json={**body, "ajuste_indice": "ipc"}).status_code == 422


def test_ajuste_trimestral_por_ipc(proyecto):
    client, p = proyecto
    for mes, v in [("2026-02-01", 2), ("2026-03-01", 3), ("2026-04-01", 4), ("2026-05-01", 5)]:
        client.post("/api/ipc", json={"mes": mes, "variacion": v})
    client.put(f"/api/proyectos/{p['id']}", json={**p, "ajuste_indice": "ipc", "ajuste_cada_meses": 3,
                                                  "ajuste_mes_base": "2026-01-01"})
    precio = lambda f: client.get(f"/api/proyectos/{p['id']}/seguimiento", params={"fecha": f}).json()  # noqa: E731
    assert precio("2026-03-15")["precio_base"] == pytest.approx(1_000_000)  # aún no venció el trimestre
    s = precio("2026-04-10")  # ajuste de abril: IPC de febrero y marzo (el de abril aún no se publicó)
    assert s["precio_base"] == pytest.approx(1_000_000 * 1.02 * 1.03)
    assert s["contrato"]["hasta"] == "2026-03-01"
    s = precio("2026-07-10")  # ajuste de julio: febrero a junio; falta junio, se usa lo publicado
    assert s["precio_base"] == pytest.approx(1_000_000 * 1.02 * 1.03 * 1.04 * 1.05)
    assert s["contrato"]["faltan_datos"] == 1


def test_factor_aumento_de_costos():
    assert contratos.factor_aumento_costos(0, date(2026, 1, 1), date(2026, 7, 1)) == 1
    f = contratos.factor_aumento_costos(0.02, date(2026, 1, 1), date(2027, 1, 1))
    assert f == pytest.approx(1.02 ** (365 / 30.44 / 2))


def test_cotizaciones_del_bcra(client, monkeypatch):
    def falso(url):
        assert "estadisticascambiarias" in url
        return {"results": {"fecha": "2026-10-23", "detalle": [
            {"codigoMoneda": "USD", "tipoCotizacion": 1541.0}, {"codigoMoneda": "EUR", "tipoCotizacion": 1724.5},
            {"codigoMoneda": "BRL", "tipoCotizacion": 284.0}]}}
    monkeypatch.setattr(bcra, "_get_json", falso)
    r = client.post("/api/cotizaciones/bcra").json()
    assert r["fecha"] == "2026-10-23" and {c["moneda"] for c in r["cotizaciones"]} == {"USD", "EUR", "BRL"}
    v = client.get("/api/cotizaciones/vigente", params={"moneda": "USD"}).json()
    assert v["valor"] == pytest.approx(1541.0)


def test_cotizaciones_cruzadas_con_base_usd(client, monkeypatch):
    cfg = client.get("/api/config").json()
    client.put("/api/config", json={**cfg, "moneda_base": "USD"})
    monkeypatch.setattr(bcra, "_get_json", lambda url: {"results": {"fecha": "2026-10-23", "detalle": [
        {"codigoMoneda": "USD", "tipoCotizacion": 1500.0}, {"codigoMoneda": "EUR", "tipoCotizacion": 1725.0}]}})
    client.post("/api/cotizaciones/bcra")
    cots = {c["moneda"]: c["valor"] for c in client.get("/api/cotizaciones").json()}
    assert cots["EUR"] == pytest.approx(1.15) and cots["ARS"] == pytest.approx(1 / 1500, rel=1e-6)


def test_ipc_del_bcra_y_errores(client, monkeypatch):
    monkeypatch.setattr(bcra, "_get_json", lambda url: {"results": [{"detalle": [
        {"fecha": "2026-09-30", "valor": 1.9}, {"fecha": "2026-08-31", "valor": 2.1}]}]})
    r = client.post("/api/ipc/bcra").json()
    assert r["meses"] == 2 and r["ultimo_mes"] == "2026-09-01"
    assert [i["variacion"] for i in client.get("/api/ipc").json()] == [1.9, 2.1]

    def sin_red(url):
        raise bcra.ErrorBCRA("No se pudo conectar con el BCRA.")
    monkeypatch.setattr(bcra, "_get_json", sin_red)
    r = client.post("/api/cotizaciones/bcra")
    assert r.status_code == 502 and "BCRA" in r.json()["detail"]


def test_presupuesto_en_dolares_fija_el_contrato(client):
    client.post("/api/cotizaciones", json={"fecha": "2026-01-01", "moneda": "USD", "valor": 1000})
    rol = client.post("/api/roles", json={"nombre": "X"}).json()
    client.post("/api/personas", json={"nombre": "A", "rol_id": rol["id"], "tipo": "autonomo", "moneda": "ARS",
                                       "tarifa_hora": 10000})
    cli = client.post("/api/clientes", json={"nombre": "C"}).json()
    p = client.post("/api/proyectos", json={"cliente_id": cli["id"], "nombre": "P", "fecha_inicio": "2026-11-02",
                                            "fecha_fin": "2026-12-11", "estado": "en_presupuesto"}).json()
    pres = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={}).json()
    d = client.put(f"/api/presupuestos/{pres['id']}", json={
        "margen": 0.5, "precio_base": 2_000_000, "moneda_contrato": "USD",
        "items": [{"tarea": "T", "rol_id": rol["id"], "horas": 10}]}).json()
    assert d["contrato"]["precio"] == pytest.approx(2000)
    client.post(f"/api/presupuestos/{pres['id']}/aprobar", json={})
    out = client.get(f"/api/proyectos/{p['id']}").json()
    assert out["moneda_contrato"] == "USD" and out["precio_contrato"] == pytest.approx(2000)


def test_documentos_en_moneda_del_contrato_y_clausula_ipc(proyecto):
    import io

    from openpyxl import load_workbook
    client, p = proyecto
    client.put(f"/api/proyectos/{p['id']}", json={**p, "moneda_contrato": "USD", "precio_contrato": 1000})
    r = client.get(f"/api/export/cliente/{p['id']}", params={"formato": "xlsx", "moneda": "ARS", "fecha": "2026-07-01"})
    texto = " ".join(str(c.value) for ws in load_workbook(io.BytesIO(r.content)) for f in ws.iter_rows()
                     for c in f if c.value is not None)
    assert "Precio (USD)" in texto and "Precio (ARS)" in texto and "1 USD = 1.500,00 ARS" in texto
    client.put(f"/api/proyectos/{p['id']}", json={**p, "ajuste_indice": "ipc", "ajuste_cada_meses": 3})
    r = client.get(f"/api/export/cliente/{p['id']}", params={"formato": "xlsx"})
    texto = " ".join(str(c.value) for ws in load_workbook(io.BytesIO(r.content)) for f in ws.iter_rows()
                     for c in f if c.value is not None)
    assert "forma trimestral por la variación del IPC (INDEC) desde diciembre 2025" in texto
