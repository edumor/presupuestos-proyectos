# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
"""Licencias, aviso de versión, respaldos e importación desde Excel."""
import io
import os
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from app import actualizaciones, ed25519, licencia

SECRETO = bytes(range(32))
PUBLICA = ed25519.clave_publica(SECRETO)


def test_ed25519_vector_rfc8032():
    sk = bytes.fromhex("9d61b19deffd5a60ba844af492ec2cc44449c5697b326919703bac031cae7f60")
    assert ed25519.clave_publica(sk).hex() == "d75a980182b10ab7d54bfed3c964073a0ee172f3daa62325af021a68f707511a"
    assert ed25519.firmar(sk, b"").hex().startswith("e5564300c360ac729086e2cc806e828a")


def test_validar_licencia():
    clave = licencia.emitir(SECRETO, "Agencia X S.R.L.", "30-1", "2026-0001", date(2027, 1, 31))
    r = licencia.validar(clave, date(2026, 12, 1), PUBLICA)
    assert r["estado"] == "registrada" and r["titular"] == "Agencia X S.R.L." and r["leyenda"].endswith("S.R.L.")
    assert licencia.validar(clave, date(2027, 2, 1), PUBLICA)["estado"] == "vencida"
    # con espacios y saltos de línea en el medio (pegada desde un mail) sigue valiendo
    assert licencia.validar(clave[:40] + "\n " + clave[40:], date(2026, 12, 1), PUBLICA)["estado"] == "registrada"
    adulterada = clave.replace(clave.split(".")[0], licencia.PREFIJO + licencia._b64(b'{"t":"Otro"}'))
    assert licencia.validar(adulterada, publica=PUBLICA)["estado"] == "invalida"
    assert licencia.validar("cualquier cosa")["estado"] == "invalida"
    assert licencia.validar(None)["estado"] == "sin_licencia"
    # firmada con otra clave: no vale para el programa
    assert licencia.validar(clave)["estado"] == "invalida"


def test_activar_licencia_api(client, monkeypatch):
    monkeypatch.setattr(licencia, "CLAVE_PUBLICA", PUBLICA)
    monkeypatch.setattr(licencia.validar, "__defaults__", (None, PUBLICA))
    assert client.get("/api/licencia").json()["estado"] == "sin_licencia"
    assert client.get("/api/config").json()["licencia_info"]["leyenda"] == "Versión sin registrar"
    assert client.put("/api/licencia", json={"clave": "PRES1-xx.yy"}).status_code == 422
    clave = licencia.emitir(SECRETO, "Agencia X")
    r = client.put("/api/licencia", json={"clave": clave}).json()
    assert r["estado"] == "registrada"
    assert client.get("/api/config").json()["licencia_info"]["titular"] == "Agencia X"
    # guardar la configuración no borra la licencia
    cfg = client.get("/api/config").json()
    client.put("/api/config", json=cfg)
    assert client.get("/api/licencia").json()["estado"] == "registrada"
    assert client.delete("/api/licencia").status_code == 204
    assert client.get("/api/licencia").json()["estado"] == "sin_licencia"


def test_aviso_de_version(client, monkeypatch):
    monkeypatch.setattr(actualizaciones, "_cache", {})
    monkeypatch.setattr(actualizaciones, "_get_json", lambda url: {
        "tag_name": "v99.1.0", "name": "Presupuestos 99.1.0", "html_url": "https://github.com/x/releases/tag/v99.1.0",
        "published_at": "2027-01-10T12:00:00Z", "body": "Novedades",
        "assets": [{"name": "Presupuestos-Setup-99.1.0.exe", "browser_download_url": "https://x/setup.exe"}]})
    r = client.get("/api/actualizacion").json()
    assert r["hay_nueva"] and r["ultima"] == "99.1.0" and r["instalador"] == "https://x/setup.exe"
    assert actualizaciones.version_tupla("0.10.0") > actualizaciones.version_tupla("0.9.3")
    cfg = client.get("/api/config").json()
    client.put("/api/config", json={**cfg, "buscar_actualizaciones": False})
    assert client.get("/api/actualizacion").json().get("desactivado")

    def sin_red(url):
        raise OSError("sin internet")
    monkeypatch.setattr(actualizaciones, "_get_json", sin_red)
    r = client.get("/api/actualizacion", params={"forzar": True}).json()
    assert r["ok"] is False and r["hay_nueva"] is False


def test_descargar_y_restaurar_respaldo(client):
    client.post("/api/clientes", json={"nombre": "Antes"})
    copia = client.get("/api/respaldos/descargar")
    assert copia.status_code == 200 and copia.content[:15] == b"SQLite format 3"
    client.post("/api/clientes", json={"nombre": "Después"})
    assert len(client.get("/api/clientes").json()) == 2
    r = client.post("/api/respaldos/restaurar", content=copia.content,
                    headers={"Content-Type": "application/octet-stream"}).json()
    assert r["ok"] and r["copia_previa"].startswith("previa_restauracion_")
    assert [c["nombre"] for c in client.get("/api/clientes").json()] == ["Antes"]
    # la copia previa aparece en la lista y se puede volver a ella
    lista = client.get("/api/respaldos").json()
    assert lista[0]["nombre"] == r["copia_previa"] and lista[0]["antes_de_restaurar"]
    client.post(f"/api/respaldos/{r['copia_previa']}/restaurar")
    assert len(client.get("/api/clientes").json()) == 2
    # archivos que no son respaldos
    malo = client.post("/api/respaldos/restaurar", content=b"hola", headers={"Content-Type": "application/octet-stream"})
    assert malo.status_code == 422
    assert client.post("/api/respaldos/noexiste.db/restaurar").status_code == 404


def _libro(hojas: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, url, contenido):
    return client.post(url, content=contenido, headers={"Content-Type": "application/octet-stream"})


def test_plantilla_de_importacion(client):
    r = client.get("/api/importar/plantilla")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Instrucciones", "Equipo", "Gastos fijos", "Servicios", "Clientes", "Horas", "Facturas", "Cobros"]
    assert wb["Equipo"]["A1"].value == "Nombre *"


def test_importar_desde_excel(client):
    client.post("/api/cotizaciones", json={"fecha": "2026-01-01", "moneda": "USD", "valor": 1000})
    contenido = _libro({
        "Equipo": [["Nombre *", "Rol *", "Tipo *", "Moneda", "Sueldo bruto anual", "Tarifa por hora", "Contribuciones %"],
                   ["Ana", "Diseño", "Empleado", None, "26.000.000", None, 28.4],
                   ["Beto", "Desarrollo", "Autónomo", "USD", None, 20, None]],
        "Gastos fijos": [["Concepto *", "Periodicidad *", "Moneda", "Importe *"], ["Alquiler", "Mensual", None, 900000]],
        "Servicios": [["Nombre *", "Rol sugerido", "Horas sugeridas", "% sobre el total", "Etapa sugerida"],
                      ["Landing", "Desarrollo", 16, None, "Ejecución"], ["Gestión", None, None, 10, None]],
        "Clientes": [["Nombre *", "País", "Moneda a mostrar"], ["Cliente Uno", "Argentina", "USD"]],
    })
    prev = _post(client, "/api/importar/vista-previa", contenido).json()
    assert prev["errores"] == 0 and not prev["importado"] and prev["total"] == 6
    assert client.get("/api/personas").json() == []  # la vista previa no guarda nada
    r = _post(client, "/api/importar", contenido).json()
    assert r["importado"]
    personas = {p["nombre"]: p for p in client.get("/api/personas").json()}
    assert personas["Ana"]["sueldo_bruto_anual"] == 26_000_000 and personas["Ana"]["ss_empresa"] == pytest.approx(0.284)
    assert personas["Beto"]["moneda"] == "USD" and personas["Beto"]["rol_nombre"] == "Desarrollo"
    srv = {s["nombre"]: s for s in client.get("/api/servicios").json()}
    assert srv["Gestión"]["pct_sobre_total"] == pytest.approx(0.10)
    # segunda vez: se actualiza, no se duplica
    r = _post(client, "/api/importar", contenido).json()
    assert r["hojas"][0]["actualizados"] == 2 and len(client.get("/api/personas").json()) == 2

    # horas, facturas y cobros sobre un proyecto existente
    cli = client.get("/api/clientes").json()[0]
    p = client.post("/api/proyectos", json={"cliente_id": cli["id"], "nombre": "Web", "fecha_inicio": "2026-01-05",
                                            "fecha_fin": "2026-03-31", "precio_base": 5_000_000,
                                            "usar_plantilla": True}).json()
    client.post(f"/api/proyectos/{p['id']}/tareas", json={"nombre": "Diseño", "persona_id": personas["Ana"]["id"],
                                                          "fecha_inicio": "2026-01-05", "fecha_fin": "2026-02-27",
                                                          "horas_previstas": 40})
    contenido = _libro({
        "Horas": [["Fecha *", "Persona *", "Proyecto *", "Tarea *", "Horas *"],
                  [date(2026, 1, 12), "Ana", "Web", "Diseño", 6], ["13/01/2026", "ana", "web", "diseño", "4,5"],
                  ["14/01/2026", "Zoe", "Web", "Diseño", 2], ["15/01/2026", "Ana", "Web", "Nada", 2]],
        "Facturas": [["Proyecto *", "Número *", "Fecha *", "Vencimiento", "Importe neto *", "Impuestos", "Etapa"],
                     ["Web", "A-0001-00000001", "10/01/2026", "20/01/2026", 2_000_000, 420_000, "Inicio"]],
        "Cobros": [["Proyecto *", "Número de factura *", "Fecha *", "Importe *"],
                   ["Web", "A-0001-00000001", "18/01/2026", 1_000_000]],
    })
    prev = _post(client, "/api/importar/vista-previa", contenido).json()
    assert prev["errores"] == 2 and not prev["importado"]
    horas = next(h for h in prev["hojas"] if h["hoja"] == "Horas")
    assert [e["fila"] for e in horas["errores"]] == [4, 5] and "Zoe" in horas["errores"][0]["mensaje"]
    assert _post(client, "/api/importar", contenido).status_code == 422
    assert client.get("/api/horas").json() == []  # con errores no se importa nada

    wb = load_workbook(io.BytesIO(contenido))
    wb["Horas"].delete_rows(4, 2)
    buf = io.BytesIO(); wb.save(buf)
    r = _post(client, "/api/importar", buf.getvalue()).json()
    assert r["importado"] and r["total"] == 4
    assert sorted(h["horas"] for h in client.get("/api/horas").json()) == [4.5, 6]
    f = client.get(f"/api/proyectos/{p['id']}/facturas").json()["facturas"][0]
    assert f["etapa_nombre"] == "Inicio" and f["cobrado"] == 1_000_000 and f["saldo"] == 1_420_000
    # reimportar: horas y cobros iguales se omiten
    r = _post(client, "/api/importar/vista-previa", buf.getvalue()).json()
    assert {h["hoja"]: h["omitidos"] for h in r["hojas"]} == {"Horas": 2, "Facturas": 0, "Cobros": 1}

    assert _post(client, "/api/importar/vista-previa", b"no es excel").status_code == 422
    assert _post(client, "/api/importar/vista-previa", _libro({"Otra": [["a"]]})).status_code == 422


def test_sin_clave_privada_en_el_repositorio():
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(raiz, ".gitignore"), encoding="utf-8") as f:
        assert "licencias_privado/" in f.read()
