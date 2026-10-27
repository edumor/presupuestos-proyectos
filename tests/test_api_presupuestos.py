# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
import io
import re

import pytest
from openpyxl import load_workbook


@pytest.fixture()
def base(client):
    """Un rol con una persona autónoma (ARS 10.000/h, sin estructura) y un proyecto sin tareas."""
    rol = client.post("/api/roles", json={"nombre": "Diseño"}).json()
    per = client.post("/api/personas", json={"nombre": "Pia", "rol_id": rol["id"], "tipo": "autonomo",
                                              "moneda": "ARS", "tarifa_hora": 10000}).json()
    cfg = client.get("/api/config").json()
    cfg["horas_autonomos_anuales"] = 0
    client.put("/api/config", json=cfg)
    cli = client.post("/api/clientes", json={"nombre": "Cliente"}).json()
    p = client.post("/api/proyectos", json={"cliente_id": cli["id"], "nombre": "Web", "fecha_inicio": "2026-11-02",
                                            "fecha_fin": "2026-12-11", "estado": "en_presupuesto"}).json()
    return client, rol, per, p


def _items(rol, per):
    return [
        {"tarea": "Diseño de piezas", "etapa": "Ejecución", "rol_id": rol["id"], "horas": 20},
        {"tarea": "Brief", "etapa": "Inicio", "persona_id": per["id"], "horas": 10},
        {"tarea": "Gestión", "etapa": "Ejecución", "rol_id": rol["id"], "pct_sobre_total": 0.10},
        {"tarea": "Consultoría externa", "etapa": "Planificación", "horas": 5, "costo_hora_base": 20000},
    ]


def test_calculo_de_un_presupuesto(base):
    client, rol, per, p = base
    pres = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={}).json()
    assert pres["version"] == 1 and pres["estado"] == "borrador"
    d = client.put(f"/api/presupuestos/{pres['id']}", json={
        "margen": 0.5, "impuesto": 0.21, "items": _items(rol, per),
        "externos": [{"concepto": "Licencia", "moneda": "ARS", "importe": 50000}]}).json()
    horas = {i["tarea"]: i["horas"] for i in d["items"]}
    assert horas["Gestión"] == pytest.approx(3.5)  # 10 % de las 35 horas fijas
    r = d["resumen"]
    # (20 + 10 + 3,5) h × 10.000 + 5 h × 20.000 + 50.000 de externos
    assert r["costo_total"] == pytest.approx(335000 + 100000 + 50000)
    assert r["precio"] == pytest.approx(485000 / 0.5)
    assert r["impuesto"] == pytest.approx(970000 * 0.21)
    # precio final manual
    d = client.put(f"/api/presupuestos/{pres['id']}", json={"margen": 0.5, "precio_base": 1_000_000,
                                                           "items": _items(rol, per)}).json()
    assert d["resumen"]["precio"] == 1_000_000 and d["resumen"]["margen_real"] == pytest.approx(1 - 435000 / 1e6)


def test_versiones_estados_y_aprobacion(base):
    client, rol, per, p = base
    v1 = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={}).json()
    client.put(f"/api/presupuestos/{v1['id']}", json={"margen": 0.35, "items": _items(rol, per)})
    assert client.post(f"/api/presupuestos/{v1['id']}/estado", json={"estado": "enviado"}).json()["estado"] == "enviado"
    # una versión enviada no se edita: se crea otra a partir de ella
    assert client.put(f"/api/presupuestos/{v1['id']}", json={"margen": 0.4}).status_code == 409
    v2 = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={"copiar_de": v1["id"]}).json()
    assert v2["version"] == 2 and len(v2["items"]) == 4

    ap = client.post(f"/api/presupuestos/{v2['id']}/aprobar", json={})
    assert ap.status_code == 200, ap.text
    proy = client.get(f"/api/proyectos/{p['id']}").json()
    assert proy["precio_base"] == pytest.approx(ap.json()["resumen"]["precio"])
    assert proy["estado"] == "en_curso"
    estados = {x["version"]: x["estado"] for x in client.get(f"/api/proyectos/{p['id']}/presupuestos").json()}
    assert estados == {1: "rechazado", 2: "aprobado"}

    tareas = client.get(f"/api/proyectos/{p['id']}/tareas").json()
    etapas = {e["id"]: e["nombre"] for e in client.get(f"/api/proyectos/{p['id']}/etapas").json()}
    assert {t["nombre"]: etapas[t["etapa_id"]] for t in tareas}["Brief"] == "Inicio"
    assert sum(t["costo_previsto_base"] for t in tareas) == pytest.approx(435000)
    # costos congelados: cambiar la tarifa no cambia la versión aprobada
    client.put(f"/api/personas/{per['id']}", json={**per, "tarifa_hora": 99999})
    assert client.get(f"/api/presupuestos/{v2['id']}").json()["resumen"]["costo_total"] == pytest.approx(435000)
    # no se aprueba otra ni se borra la aprobada
    v3 = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={"copiar_de": v2["id"]}).json()
    assert client.post(f"/api/presupuestos/{v3['id']}/aprobar", json={}).status_code == 409
    assert client.delete(f"/api/presupuestos/{v2['id']}").status_code == 409
    assert client.delete(f"/api/presupuestos/{v3['id']}").status_code == 204
    # el proyecto con todo se puede borrar
    assert client.delete(f"/api/proyectos/{p['id']}").status_code == 204


def test_item_sin_costo_impide_aprobar(base):
    client, rol, per, p = base
    pres = client.post(f"/api/proyectos/{p['id']}/presupuestos", json={}).json()
    d = client.put(f"/api/presupuestos/{pres['id']}", json={"margen": 0.3, "items": [{"tarea": "X", "horas": 5}]}).json()
    assert d["avisos"]
    assert client.post(f"/api/presupuestos/{pres['id']}/aprobar", json={}).status_code == 409


def test_propuesta_desde_el_presupuesto(client):
    client.post("/api/ejemplo")
    pid = next(p["id"] for p in client.get("/api/proyectos").json() if p["nombre"] == "Rediseño del sitio web")
    pres = client.get(f"/api/proyectos/{pid}/presupuestos").json()[0]
    assert pres["estado"] == "enviado" and pres["n_items"] == 8
    r = client.get(f"/api/export/presupuesto/{pres['id']}", params={"formato": "xlsx", "moneda": "USD"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    texto = " ".join(str(c.value) for ws in wb for f in ws.iter_rows() for c in f if c.value is not None)
    palabras = set(re.findall(r"\w+", texto.lower()))
    assert not {"costo", "margen", "tarifa"} & palabras
    filas = list(wb["Propuesta"].iter_rows(values_only=True))
    total = next(f for f in filas if f[0] == "Total")
    assert total[5] == pytest.approx(pres["precio"]) and total[6] == pytest.approx(pres["precio"] / 1540, abs=0.01)
    pdf = client.get(f"/api/export/presupuesto/{pres['id']}", params={"idioma": "en"})
    assert pdf.content[:4] == b"%PDF"


def test_catalogo_de_servicios(client):
    rol = client.post("/api/roles", json={"nombre": "PM"}).json()
    s = client.post("/api/servicios", json={"nombre": "Gestión", "rol_sugerido_id": rol["id"], "pct_sobre_total": 0.1,
                                            "etapa_sugerida": "Ejecución"}).json()
    assert s["rol_nombre"] == "PM"
    assert client.post("/api/servicios", json={"nombre": "Sin horas"}).status_code == 422
    assert client.delete(f"/api/servicios/{s['id']}").status_code == 204
    assert client.get("/api/servicios").json() == []
