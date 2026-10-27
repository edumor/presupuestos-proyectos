# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (c) 2026 Eduardo José Moreno. Desarrollo e Implementación: Lic. Eduardo J. Moreno
import base64
import io
import json
import re

import pytest
from openpyxl import load_workbook

PERSONAS_EJEMPLO = ["Ana", "Bruno", "Carla", "Diego", "Elena", "Fede", "Gina"]
PALABRAS_INTERNAS = ["costo", "costos", "cost", "costs", "margin", "margen", "cpi", "spi", "rate", "tarifa", "markup"]


def _textos_xlsx(contenido: bytes) -> str:
    wb = load_workbook(io.BytesIO(contenido))
    return " ".join(str(c.value) for ws in wb for fila in ws.iter_rows() for c in fila if c.value is not None)


@pytest.fixture()
def ejemplo(client):
    client.post("/api/ejemplo")
    pid = next(p["id"] for p in client.get("/api/proyectos").json() if p["nombre"] == "Campaña de lanzamiento")
    return client, pid


@pytest.mark.parametrize("tipo", ["propuesta", "avance"])
@pytest.mark.parametrize("formato,firma", [("pdf", b"%PDF"), ("xlsx", b"PK")])
def test_documentos_cliente_se_generan(ejemplo, tipo, formato, firma):
    client, pid = ejemplo
    r = client.get(f"/api/export/cliente/{pid}", params={"tipo": tipo, "formato": formato})
    assert r.status_code == 200 and r.content[:len(firma)] == firma
    assert r.headers["content-disposition"].startswith("attachment;")


@pytest.mark.parametrize("idioma", ["es", "en"])
@pytest.mark.parametrize("tipo", ["propuesta", "avance"])
def test_excel_cliente_sin_datos_internos(ejemplo, tipo, idioma):
    client, pid = ejemplo
    texto = _textos_xlsx(client.get(f"/api/export/cliente/{pid}", params={"tipo": tipo, "formato": "xlsx",
                                                                          "idioma": idioma}).content)
    for nombre in PERSONAS_EJEMPLO:
        assert nombre not in texto.split(), f"aparece el nombre {nombre}"
    palabras = set(re.findall(r"\w+", texto.lower()))
    for palabra in PALABRAS_INTERNAS:
        assert palabra not in palabras, f"aparece '{palabra}'"
    assert "Desarrollo e Implementación: Lic. Eduardo J. Moreno" in texto


def _precios(contenido: bytes, hoja: str, etapa: str, col: int = 5):
    ws = load_workbook(io.BytesIO(contenido))[hoja]
    filas = list(ws.iter_rows(values_only=True))
    i = next(k for k, f in enumerate(filas) if f and f[0] == etapa)
    precios = []
    for f in filas[i + 1:]:
        if f[0] == "Total":
            return precios, f[col], filas[i]
        if f[0]:
            precios.append(f[col])


@pytest.mark.parametrize("idioma,hoja,etapa", [("es", "Propuesta", "Etapa"), ("en", "Proposal", "Stage")])
def test_precio_por_etapa_suma_el_total(ejemplo, idioma, hoja, etapa):
    client, pid = ejemplo
    precio = client.get(f"/api/proyectos/{pid}").json()["precio_base"]
    contenido = client.get(f"/api/export/cliente/{pid}", params={"tipo": "propuesta", "formato": "xlsx",
                                                                 "idioma": idioma}).content
    precios, total, cab = _precios(contenido, hoja, etapa)
    assert "ARS" in cab[5]
    assert total == pytest.approx(precio)
    assert sum(precios) == pytest.approx(total, abs=0.001)


def test_propuesta_convertida_a_dolares(ejemplo):
    client, pid = ejemplo
    precio = client.get(f"/api/proyectos/{pid}").json()["precio_base"]
    contenido = client.get(f"/api/export/cliente/{pid}", params={"tipo": "propuesta", "formato": "xlsx",
                                                                 "moneda": "USD"}).content
    _, total_usd, cab = _precios(contenido, "Propuesta", "Etapa", col=6)
    assert "USD" in cab[6]
    assert total_usd == pytest.approx(precio / 1540, abs=0.01)  # 1 USD = 1.540 ARS en el ejemplo
    pdf = client.get(f"/api/export/cliente/{pid}", params={"tipo": "propuesta", "moneda": "USD", "idioma": "en"})
    assert pdf.status_code == 200 and pdf.headers["content-disposition"].count("Proposal")


def test_idioma_de_los_documentos(ejemplo):
    client, pid = ejemplo
    es = _textos_xlsx(client.get(f"/api/export/cliente/{pid}", params={"tipo": "avance", "formato": "xlsx"}).content)
    en = _textos_xlsx(client.get(f"/api/export/cliente/{pid}", params={"tipo": "avance", "formato": "xlsx",
                                                                       "idioma": "en"}).content)
    assert "Informe de avance del proyecto" in es and "Project Progress Report" in en
    html_es = client.get(f"/api/export/dashboard/{pid}", params={"version": "cliente"}).text
    html_en = client.get(f"/api/export/dashboard/{pid}", params={"version": "cliente", "idioma": "en"}).text
    assert '<html lang="es">' in html_es and '<html lang="en">' in html_en


def test_moneda_sin_cotizacion_da_422(client):
    client.post("/api/ejemplo")
    pid = client.get("/api/proyectos").json()[0]["id"]
    client.delete("/api/cotizaciones/%d" % next(c["id"] for c in client.get("/api/cotizaciones").json()
                                                  if c["moneda"] == "GBP"))
    for c in [c for c in client.get("/api/cotizaciones").json() if c["moneda"] == "GBP"]:
        client.delete(f"/api/cotizaciones/{c['id']}")
    r = client.get(f"/api/export/cliente/{pid}", params={"tipo": "propuesta", "moneda": "GBP"})
    assert r.status_code == 422 and r.json()["moneda"] == "GBP"


def test_dashboard_cliente_sin_datos_internos(ejemplo):
    client, pid = ejemplo
    html = client.get(f"/api/export/dashboard/{pid}", params={"version": "cliente"}).text
    datos = json.loads(html.split("const D = ", 1)[1].split(";\nconst CLI", 1)[0])
    assert datos["logo"] is None  # sin logo cargado no lleva logo
    plano = json.dumps({k: v for k, v in datos.items() if k != "logo"}, ensure_ascii=False)
    for nombre in PERSONAS_EJEMPLO:
        assert f'"{nombre}"' not in plano
    for clave in ["costo", "precio", "usd", "cpi", "spi", "margen", "persona"]:
        assert clave not in plano.lower()
    assert datos["credito"] == "Desarrollo e Implementación: Lic. Eduardo J. Moreno"
    assert "<script>" in html and "Chart" in html  # Chart.js embebido: funciona sin internet


def test_dashboard_interno(ejemplo):
    client, pid = ejemplo
    r = client.get(f"/api/export/dashboard/{pid}", params={"version": "interno"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    datos = json.loads(r.text.split("const D = ", 1)[1].split(";\nconst CLI", 1)[0])
    assert datos["m"]["cpi"] is not None and datos["tareas"] and datos["por_persona"]


def test_datos_empresa_y_logo(client):
    cfg = client.get("/api/config").json()
    cfg.update(empresa_nombre="Consultora X", validez_dias=45, condiciones="Net 30")
    out = client.put("/api/config", json=cfg).json()
    assert out["empresa_nombre"] == "Consultora X" and out["validez_dias"] == 45 and not out["tiene_logo"]
    # sin logo cargado no hay logo (cada empresa carga el suyo)
    assert client.get("/api/config/logo").status_code == 404
    assert client.put("/api/config/logo", json={"data_url": "data:image/png;base64,bm8gZXMgdW5hIGltYWdlbg=="}).status_code == 422
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "blue").save(buf, "PNG")
    url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    assert client.put("/api/config/logo", json={"data_url": url}).status_code == 204
    assert client.get("/api/config").json()["tiene_logo"]
    assert client.get("/api/config/logo").headers["content-type"] == "image/png"
    client.delete("/api/config/logo")
    assert not client.get("/api/config").json()["tiene_logo"]


def _subir_logo(client):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "blue").save(buf, "PNG")
    url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    assert client.put("/api/config/logo", json={"data_url": url}).status_code == 204


def test_logo_solo_en_la_barra_de_los_html(ejemplo):
    client, pid = ejemplo
    _subir_logo(client)
    for v in ("cliente", "interno"):
        datos = json.loads(client.get(f"/api/export/dashboard/{pid}", params={"version": v}).text
                           .split("const D = ", 1)[1].split(";\nconst ", 1)[0])
        assert datos["logo"].startswith("data:image/png;base64,")
        html = client.get(f"/api/export/dashboard/{pid}", params={"version": v}).text
        cab = html.split("<header>", 1)[1].split("</header>", 1)[0]
        assert cab.index('id="logo"') < cab.index('id="titulo"')  # a la izquierda del título
        assert "header { background:#fff;" in html  # barra blanca
    wb = load_workbook(io.BytesIO(client.get(f"/api/export/cliente/{pid}", params={"formato": "xlsx"}).content))
    assert not wb.worksheets[0]._images  # el Excel no lleva logo
    pdf = client.get(f"/api/export/cliente/{pid}", params={"formato": "pdf"}).content
    assert b"/Subtype /Image" not in pdf  # el PDF no lleva logo
